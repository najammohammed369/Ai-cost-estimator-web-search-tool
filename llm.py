from openai.types.responses import response
from pdf_parser import  PDFExtractor
from openai import OpenAI
import re 
import pandas as pd
import json
from openpyxl.styles import Alignment, Font

class LLM():                                                                                                                                                           
    def __init__(self):
        self.client = OpenAI(
            base_url=input("Enter model endpoint url: "),
            api_key=input("enter api key: ")
        )

    def jsonify_data(self):         
        text = PDFExtractor.pdf_text_extractor("C:\\Users\\najam\\CloudMetica\\Ai-cost-estimator-web-search-tool\\VANTAGE-OUTLINE-SPECIFICATION-SGD_FEB2020_rev-02.pdf")
        text_str = "\n".join(text) if isinstance(text, list) else text
        response = self.client.chat.completions.create(
        model="databricks-meta-llama-3-3-70b-instruct",
        messages= [
            {"role": "system", "content": "You are a helpful assistant.go through the text provided and generate a json data of the specifications of the ship. do not hallucinate or produce any false data which is not present in the text. provide output in pure json format, do not write any heading like here is your json or this is the json of the ship specification and do not use any backticks or anything just provide plain json as output"},
            {"role": "user", "content": text_str}
        ]
    )   
        return response.choices[0].message.content

    def create_excel(self):
        data = self.jsonify_data()
        try:
            # Parse the JSON text returned by the LLM
            parsed_data = json.loads(data)
            file_path = "C:\\Users\\najam\\CloudMetica\\Ai-cost-estimator-web-search-tool\\ship_data.xlsx"
            # Create a Pandas Excel writer
            with pd.ExcelWriter(file_path, engine='openpyxl') as writer:
                if isinstance(parsed_data, dict):
                    workbook = writer.book
                    sheet_name = "Ship Data"
                    sheet = workbook.create_sheet(sheet_name)
                    writer.sheets[sheet_name] = sheet
                    current_row = 1
                    
                    def extract_sections(node, prefix=""):
                        sections_list = []
                        scalars = {}
                        for k, v in node.items():
                            if isinstance(v, dict):
                                sections_list.extend(extract_sections(v, k))
                            elif isinstance(v, list):
                                sections_list.append((k, v, False))
                            else:
                                scalars[k] = v
                        if scalars:
                            heading = prefix if prefix else "General Specification"
                            sections_list.insert(0, (heading, scalars, True))
                        return sections_list
                        
                    sections = extract_sections(parsed_data)
                    
                    for key, value, is_dict in sections:
                        if is_dict:
                            df = pd.json_normalize(value).T.reset_index()
                        else:
                            df = pd.json_normalize(value)
                        
                        # Write the dataframe directly (no heading)
                        if is_dict:
                            df.to_excel(writer, sheet_name=sheet_name, startrow=current_row - 1, header=False, index=False)
                            current_row += len(df)
                        else:
                            df.to_excel(writer, sheet_name=sheet_name, startrow=current_row - 1, index=False)
                            current_row += len(df) + 1
                    
                # Set column width and wrap text for all cells
                from openpyxl.utils import get_column_letter
                for col in range(1, sheet.max_column + 1):
                    sheet.column_dimensions[get_column_letter(col)].width = 20
                    
                for row in sheet.iter_rows():
                    for cell in row:
                        if cell.alignment:
                            cell.alignment = Alignment(
                                horizontal=cell.alignment.horizontal,
                                vertical=cell.alignment.vertical,
                                wrap_text=True
                            )
                        else:
                            cell.alignment = Alignment(wrap_text=True)
                            
                # Remove default sheet created by openpyxl if it exists
                if "Sheet" in workbook.sheetnames:
                    del workbook["Sheet"]
                elif isinstance(parsed_data, list):
                    df = pd.DataFrame(parsed_data)
                    df.to_excel(writer, sheet_name="Ship Data", index=False)
                else:
                    df = pd.DataFrame([parsed_data])
                    df.to_excel(writer, sheet_name="Ship Data", index=False)
                
            return f"Successfully created Excel file at: {file_path}"
        except Exception as e:
            return f"Error creating Excel file: {str(e)}"